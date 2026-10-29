#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build a minimal, decoration-free paper-reading deck from a JSON spec.

Design contract (see references/style-spec.md):
  * 16:9, pure white background, black text, restrained accent colors
    (default red FF0000; optional blue 0070C0 as a second accent, max 2 accents)
  * one 40pt title per slide, body 24pt (auto-shrinks to fit, never below 14pt)
  * figures are the paper's own figures, cropped losslessly from the PDF
  * no theme shapes, no icons, no page numbers, no agenda/thanks slides

Usage:
  python build_deck.py deck.json -o out.pptx
  python build_deck.py deck.json -o out.pptx --report
"""
import argparse
import json
import math
import re
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.oxml.ns import qn
except ImportError:
    print("ERROR: python-pptx is required. Install with: pip install python-pptx")
    sys.exit(1)

# ---------------------------------------------------------------- constants
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN_L = 0.92
MARGIN_R = 0.40
TITLE_TOP = 0.50
TITLE_SIZE = 40
BODY_TOP = 1.72
BODY_BOTTOM = 7.02
LINE_FACTOR = 1.30          # line height / font size
FONT_ZH = "等线"
FONT_EN = "等线"

BLACK = RGBColor(0x00, 0x00, 0x00)
RED = RGBColor(0xFF, 0x00, 0x00)
BLUE = RGBColor(0x00, 0x70, 0xC0)
GREY = RGBColor(0x59, 0x59, 0x59)

SIZE_LADDER = [26, 24, 22, 20, 18, 17, 16, 15, 14]

TOKEN_RE = re.compile(r"(<rb>.*?</rb>|<bb>.*?</bb>|<r>.*?</r>|<b>.*?</b>|\*\*.*?\*\*)", re.S)
EA_SUCCESSORS = ("a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")


# ---------------------------------------------------------------- helpers
def parse_segments(text):
    """Split inline markup into (text, bold, color) runs.

    Supported: **bold**, <r>red</r>, <rb>red+bold</rb>, <b>blue</b>, <bb>blue+bold</bb>.
    color is an RGBColor or None (inherit from the paragraph's base color).
    Nested **bold** inside a color tag is honoured.
    """
    segs = []
    for tok in TOKEN_RE.split(text):
        if not tok:
            continue
        if tok.startswith("<rb>"):
            segs.append((tok[4:-5], True, RED))
        elif tok.startswith("<bb>"):
            segs.append((tok[4:-5], True, BLUE))
        elif tok.startswith("<r>"):
            inner = tok[3:-4]
            bold = inner.startswith("**") and inner.endswith("**") and len(inner) > 4
            segs.append((inner[2:-2] if bold else inner, bold, RED))
        elif tok.startswith("<b>"):
            inner = tok[3:-4]
            bold = inner.startswith("**") and inner.endswith("**") and len(inner) > 4
            segs.append((inner[2:-2] if bold else inner, bold, BLUE))
        elif tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            segs.append((tok[2:-2], True, None))
        else:
            segs.append((tok, False, None))
    return segs or [("", False, None)]


def plain(text):
    return "".join(s[0] for s in parse_segments(text))


def text_units(text):
    """Approximate width of a string in em units (CJK = 1.0, latin = 0.55)."""
    u = 0.0
    for ch in text:
        if ch == "\t":
            u += 1.0
        elif ord(ch) > 0x2E7F:
            u += 1.0
        elif ch in "iljtfrI.,:;'|!()[]":
            u += 0.32
        elif ch in "mwMW@":
            u += 0.85
        else:
            u += 0.55
    return u


def wrap_lines(text, width_pt, size_pt):
    per_line = max(1.0, width_pt / size_pt)
    total = 0
    for para in str(text).split("\n"):
        total += max(1, math.ceil(text_units(para) / per_line))
    return total


def _set_ea(run, name):
    rPr = run._r.get_or_add_rPr()
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.insert_element_before(ea, *EA_SUCCESSORS)
    ea.set("typeface", name)


def style_run(run, size, bold=False, color=BLACK, font=None):
    f = run.font
    f.name = font or FONT_EN
    f.size = Pt(size)
    f.bold = bold
    f.color.rgb = color
    _set_ea(run, FONT_ZH)


def add_textbox(slide, left, top, width, height, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return box


def set_bullet(para, char="•", indent_pt=26):
    pPr = para._p.get_or_add_pPr()
    pPr.set("marL", str(int(Pt(indent_pt))))
    pPr.set("indent", str(int(-Pt(indent_pt))))
    buFont = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"})
    buChar = pPr.makeelement(qn("a:buChar"), {"char": char})
    pPr.append(buFont)
    pPr.append(buChar)


def clear_bullet(para):
    pPr = para._p.get_or_add_pPr()
    pPr.set("marL", "0")
    pPr.set("indent", "0")


def add_para(tf, text, size, bullet=False, space_before=0, space_after=6,
             bold=False, color=BLACK, align=PP_ALIGN.LEFT, first=False,
             line_spacing=1.0):
    para = tf.paragraphs[0] if first else tf.add_paragraph()
    para.alignment = align
    para.line_spacing = line_spacing
    para.space_before = Pt(space_before)
    para.space_after = Pt(space_after)
    if bullet:
        set_bullet(para)
    else:
        clear_bullet(para)
    for txt, b, c in parse_segments(text):
        run = para.add_run()
        run.text = txt
        style_run(run, size, bold or b, c if c else color)
    return para


def block_height(blocks, width_in, size, gap_scale=1.0):
    width_pt = width_in * 72.0
    h = 0.0
    for i, blk in enumerate(blocks):
        sz = blk.get("sz", size)
        lines = wrap_lines(plain(blk["t"]), width_pt - (26 if blk.get("b", True) else 0), sz)
        h += lines * sz * LINE_FACTOR
        h += blk.get("gap", 8 if blk.get("b", True) else 4) * gap_scale
    return h / 72.0


def fit_size(blocks, width_in, height_in, preferred=None):
    ladder = [preferred] if preferred else SIZE_LADDER
    if preferred:
        ladder = [s for s in SIZE_LADDER if s <= preferred] or [preferred]
    for size in ladder:
        if block_height(blocks, width_in, size) <= height_in:
            return size
    return ladder[-1]


def render_blocks(slide, blocks, left, top, width, height, size, anchor=MSO_ANCHOR.TOP):
    box = add_textbox(slide, left, top, width, height, anchor)
    tf = box.text_frame
    for i, blk in enumerate(blocks):
        add_para(
            tf,
            blk["t"],
            blk.get("sz", size),
            bullet=blk.get("b", True),
            space_before=blk.get("before", 0),
            space_after=blk.get("gap", 8 if blk.get("b", True) else 4),
            align=blk.get("align", PP_ALIGN.LEFT),
            first=(i == 0),
        )
    return box


def add_title(slide, text):
    width_pt = (SLIDE_W - MARGIN_L - MARGIN_R) * 72
    units = text_units(plain(text))
    size = TITLE_SIZE
    while size > 26 and units * size > width_pt:
        size -= 2
    box = add_textbox(slide, MARGIN_L, TITLE_TOP, SLIDE_W - MARGIN_L - MARGIN_R, 0.95)
    tf = box.text_frame
    add_para(tf, text, size, bullet=False, space_after=0, first=True)
    return box


def fit_picture(slide, path, left, top, width, height, align="center"):
    """Place an image inside the box, preserving aspect ratio."""
    from PIL import Image

    with Image.open(path) as im:
        iw, ih = im.size
    ratio = iw / ih
    box_ratio = width / height
    if ratio >= box_ratio:
        w = width
        h = width / ratio
    else:
        h = height
        w = height * ratio
    if align == "center":
        l = left + (width - w) / 2
    elif align == "right":
        l = left + (width - w)
    else:
        l = left
    t = top + (height - h) / 2
    return slide.shapes.add_picture(str(path), Inches(l), Inches(t), Inches(w), Inches(h)), (l, t, w, h)


def add_caption(slide, text, left, top, width, size=12):
    box = add_textbox(slide, left, top, width, 0.3)
    add_para(box.text_frame, text, size, bullet=False, space_after=0,
             color=GREY, align=PP_ALIGN.CENTER, first=True)


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


# ---------------------------------------------------------------- slide builders
def build_cover(prs, spec):
    slide = new_slide(prs)
    venue = spec.get("venue", "")
    if venue:
        box = add_textbox(slide, 5.9, 0.38, SLIDE_W - 5.9 - 0.45, 0.75)
        add_para(box.text_frame, venue, 14, bullet=False, space_after=0,
                 align=PP_ALIGN.RIGHT, color=BLACK, first=True)
    title = f"《{spec['title_en']}》"
    box = add_textbox(slide, 0.55, 2.05, SLIDE_W - 1.1, 2.05, MSO_ANCHOR.MIDDLE)
    add_para(box.text_frame, title, 38, bullet=False, space_after=0,
             align=PP_ALIGN.CENTER, first=True, line_spacing=1.1)
    if spec.get("title_zh"):
        box = add_textbox(slide, 1.30, 4.42, SLIDE_W - 2.6, 1.55, MSO_ANCHOR.TOP)
        add_para(box.text_frame, spec["title_zh"], 24, bullet=False, space_after=0,
                 align=PP_ALIGN.CENTER, first=True, line_spacing=1.15)
    if spec.get("authors"):
        box = add_textbox(slide, 1.30, 6.10, SLIDE_W - 2.6, 0.45)
        add_para(box.text_frame, spec["authors"], 14, bullet=False, space_after=0,
                 align=PP_ALIGN.CENTER, color=GREY, first=True)


def build_points(prs, slide_spec, root=None):
    slide = new_slide(prs)
    add_title(slide, slide_spec["title"])
    top = BODY_TOP
    if slide_spec.get("lead"):
        lead_size = slide_spec.get("lead_size", 34)
        lines = wrap_lines(plain(slide_spec["lead"]), (SLIDE_W - MARGIN_L - MARGIN_R) * 72, lead_size)
        h = lines * lead_size * LINE_FACTOR / 72 + 0.18
        box = add_textbox(slide, MARGIN_L, top - 0.04, SLIDE_W - MARGIN_L - MARGIN_R, h)
        add_para(box.text_frame, slide_spec["lead"], lead_size, bullet=True,
                 space_after=0, first=True)
        top += h + 0.12
    blocks = [{"t": t, "b": True} for t in slide_spec["bullets"]]
    width = SLIDE_W - MARGIN_L - MARGIN_R
    size = fit_size(blocks, width, BODY_BOTTOM - top, slide_spec.get("size"))
    render_blocks(slide, blocks, MARGIN_L, top, width, BODY_BOTTOM - top, size)
    return size


def build_columns(prs, slide_spec, root=None):
    slide = new_slide(prs)
    add_title(slide, slide_spec["title"])
    cols = slide_spec["columns"]
    n = len(cols)
    gap = 0.45
    width = (SLIDE_W - MARGIN_L - MARGIN_R - gap * (n - 1)) / n
    top = BODY_TOP
    footer = slide_spec.get("footer")
    limit = BODY_BOTTOM - (0.80 if footer else 0.02)
    ends = []
    for i, col in enumerate(cols):
        left = MARGIN_L + i * (width + gap)
        y = top
        if col.get("head"):
            head_size = col.get("head_size", 23)
            box = add_textbox(slide, left, y, width, head_size * 1.5 / 72)
            add_para(box.text_frame, col["head"], head_size, bullet=False,
                     space_after=0, bold=True, first=True)
            y += head_size * 1.5 / 72 + 0.12
        blocks = [{"t": t, "b": True, "gap": 6} for t in col["lines"]]
        tail = col.get("tail")
        if tail:
            blocks.append({"t": tail, "b": False, "gap": 2, "before": 10})
        size = fit_size(blocks, width, limit - y, col.get("size"))
        box = render_blocks(slide, blocks, left, y, width, limit - y, size)
        if tail:
            # emphasise the closing verdict of each column, unless the tail
            # already carries its own inline markup (bold or a color tag)
            if not any(c is not None or b for _, b, c in parse_segments(tail)):
                for p in box.text_frame.paragraphs[-1:]:
                    for r in p.runs:
                        r.font.color.rgb = RED
        ends.append(y + block_height(blocks, width, size))
    if footer:
        fh = 0.62
        blocks = [{"t": footer, "b": False, "gap": 0}]
        fsize = fit_size(blocks, SLIDE_W - MARGIN_L - MARGIN_R, fh,
                         slide_spec.get("footer_size", 21))
        fy = min(max(ends) + 0.30, BODY_BOTTOM - fh)
        box = add_textbox(slide, MARGIN_L, fy, SLIDE_W - MARGIN_L - MARGIN_R, fh)
        add_para(box.text_frame, footer, fsize, bullet=False, space_after=0, first=True)


def build_figure(prs, slide_spec, root):
    slide = new_slide(prs)
    add_title(slide, slide_spec["title"])
    side = slide_spec.get("side", "right")
    img_w = slide_spec.get("img_w", 8.10)
    text_w = slide_spec.get("text_w", 3.95)
    if side == "right":
        img_left, text_left = 0.28, 0.28 + img_w + 0.35
    else:
        text_left, img_left = 0.28, SLIDE_W - 0.28 - img_w
    img_top, img_h = BODY_TOP - 0.02, BODY_BOTTOM - BODY_TOP + 0.02
    pic, rect = fit_picture(slide, root / slide_spec["image"], img_left, img_top, img_w, img_h)
    if slide_spec.get("caption"):
        add_caption(slide, slide_spec["caption"], img_left, rect[1] + rect[3] + 0.04, img_w, 11)
    blocks = [{"t": t, "b": True, "gap": 7} for t in slide_spec["bullets"]]
    size = fit_size(blocks, text_w, img_h - 0.25, slide_spec.get("size", 20))
    render_blocks(slide, blocks, text_left, BODY_TOP, text_w, img_h - 0.25, size)


def build_figure_stack(prs, slide_spec, root):
    slide = new_slide(prs)
    add_title(slide, slide_spec["title"])
    images = slide_spec["images"]
    img_w = slide_spec.get("img_w", 7.45)
    text_w = slide_spec.get("text_w", 4.30)
    text_left = 0.28 + img_w + 0.30
    top, bottom = BODY_TOP - 0.02, BODY_BOTTOM
    gap = 0.10
    h = (bottom - top - gap * (len(images) - 1)) / len(images)
    for i, im in enumerate(images):
        fit_picture(slide, root / im["path"], 0.28, top + i * (h + gap), img_w, h - 0.18, "left")
        if im.get("caption"):
            add_caption(slide, im["caption"], 0.28, top + i * (h + gap) + h - 0.20, img_w, 11)
    blocks = [{"t": t, "b": True, "gap": 7} for t in slide_spec["bullets"]]
    size = fit_size(blocks, text_w, bottom - BODY_TOP, slide_spec.get("size", 18))
    render_blocks(slide, blocks, text_left, BODY_TOP, text_w, bottom - BODY_TOP, size)


def build_figure_pair(prs, slide_spec, root):
    slide = new_slide(prs)
    add_title(slide, slide_spec["title"])
    images = slide_spec["images"]
    n = len(images)
    total_w = SLIDE_W - 0.6
    gap = 0.45
    cell = (total_w - gap * (n - 1)) / n
    img_h = slide_spec.get("img_h", 3.05)
    for i, im in enumerate(images):
        left = 0.30 + i * (cell + gap)
        fit_picture(slide, root / im["path"], left, BODY_TOP - 0.02, cell, img_h)
        if im.get("caption"):
            add_caption(slide, im["caption"], left, BODY_TOP - 0.02 + img_h + 0.02, cell, 11)
    top = BODY_TOP - 0.02 + img_h + 0.32
    blocks = [{"t": t, "b": True, "gap": 8} for t in slide_spec["bullets"]]
    width = min(SLIDE_W - 2 * MARGIN_L, slide_spec.get("text_w", 11.5))
    left = MARGIN_L
    size = fit_size(blocks, width, BODY_BOTTOM - top, slide_spec.get("size", 20))
    render_blocks(slide, blocks, left, top, width, BODY_BOTTOM - top, size)


def build_two_col(prs, slide_spec, root=None):
    slide = new_slide(prs)
    add_title(slide, slide_spec["title"])
    gap = 0.55
    width = (SLIDE_W - MARGIN_L - MARGIN_R - gap) / 2
    for i, key in enumerate(("left", "right")):
        col = slide_spec.get(key)
        if not col:
            continue
        left = MARGIN_L + i * (width + gap)
        y = BODY_TOP
        if col.get("head"):
            box = add_textbox(slide, left, y, width, 0.45)
            add_para(box.text_frame, col["head"], 22, bullet=False, space_after=0,
                     bold=True, first=True)
            y += 0.50
        blocks = [{"t": t, "b": True, "gap": 6} for t in col["bullets"]]
        size = fit_size(blocks, width, BODY_BOTTOM - y, slide_spec.get("size", 20))
        render_blocks(slide, blocks, left, y, width, BODY_BOTTOM - y, size)


BUILDERS = {
    "points": build_points,
    "columns": build_columns,
    "figure": build_figure,
    "figure_stack": build_figure_stack,
    "figure_pair": build_figure_pair,
    "two_col": build_two_col,
}


def main():
    ap = argparse.ArgumentParser(description="Build a minimal paper-reading deck from JSON.")
    ap.add_argument("spec")
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--report", action="store_true", help="print per-slide layout report")
    args = ap.parse_args()

    spec_path = Path(args.spec).resolve()
    root = spec_path.parent
    spec = json.loads(spec_path.read_text(encoding="utf-8"))

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    for i, s in enumerate(spec["slides"], 1):
        kind = s.get("type", "points")
        if kind == "cover":
            build_cover(prs, {**spec, **s})
        elif kind in BUILDERS:
            BUILDERS[kind](prs, s, root)
        else:
            raise SystemExit(f"slide {i}: unknown type {kind!r}")

    out = Path(args.out).resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out))
    print(f"SAVED {out}  slides={len(prs.slides.__iter__.__self__._sldIdLst)}")


if __name__ == "__main__":
    main()
